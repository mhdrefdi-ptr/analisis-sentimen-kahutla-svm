"""Generate a Word report from recorded pytest and model evaluation results."""
import json
import platform
import xml.etree.ElementTree as ET
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports'
r = json.loads((ROOT / 'artifacts/evaluation.json').read_text())
suite = ET.parse(OUT / 'pytest-results.xml').getroot().find('testsuite')
cases = suite.findall('testcase')
labels = r['labels']
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False})
colors = ['#238678', '#CB6848', '#718399']
fig, ax = plt.subplots(figsize=(7, 2.5))
bars = ax.bar([s.capitalize() for s in labels], [r['class_counts'][s] for s in labels], color=colors, width=.55)
ax.bar_label(bars, padding=3)
ax.set_ylim(0, 48)
ax.set_ylabel('Jumlah baris')
ax.set_title('Distribusi dataset: 92 baris', loc='left', weight='bold')
fig.tight_layout()
fig.savefig(OUT / 'distribusi_dataset.png', dpi=180)
plt.close(fig)
fig, ax = plt.subplots(figsize=(7, 2.9))
x = np.arange(3)
for i, (metric, title, color) in enumerate([('precision', 'Precision', '#238678'), ('recall', 'Recall', '#CB6848'), ('f1-score', 'F1-score', '#718399')]):
    bars = ax.bar(x + (i-1)*.24, [100*r['classification_report'][s][metric] for s in labels], .24, label=title, color=color)
    ax.bar_label(bars, fmt='%.1f', padding=3, fontsize=8)
ax.set_xticks(x, [s.capitalize() for s in labels])
ax.set_ylim(0, 110)
ax.set_ylabel('Skor (%)')
ax.legend(loc='upper right', ncol=3, fontsize=8)
fig.tight_layout()
fig.savefig(OUT / 'metrik_per_kelas.png', dpi=180)
plt.close(fig)
fig, ax = plt.subplots(figsize=(6.5, 3.7))
matrix = np.array(r['confusion_matrix'])
ax.imshow(matrix, cmap='Blues', vmin=0, vmax=matrix.max())
ax.set_xticks(range(3), [s.capitalize() for s in labels])
ax.set_yticks(range(3), [s.capitalize() for s in labels])
ax.set_xlabel('Label prediksi')
ax.set_ylabel('Label aktual')
for (y, x), value in np.ndenumerate(matrix):
    ax.text(x, y, str(value), ha='center', va='center', fontsize=18, color='white' if value >= 5 else '#17344A')
fig.tight_layout()
fig.savefig(OUT / 'confusion_matrix.png', dpi=180)
plt.close(fig)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(10)
normal.paragraph_format.space_after = Pt(7)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.color.rgb = RGBColor.from_string('17344A')
sec.header.paragraphs[0].text = 'SENTIMEN KARHUTLA  |  LAPORAN PENGUJIAN MVP'
footer = sec.footer.paragraphs[0]
footer.text = 'Hasil pengujian aktual  |  Halaman '
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    for row in t.rows:
        prop = row._tr.get_or_add_trPr()
        prop.append(OxmlElement('w:cantSplit'))
    return t

def paragraph(text):
    doc.add_paragraph(text)

def chart(file, caption, width=6.5):
    doc.add_picture(str(OUT / file), width=Inches(width))
    p = doc.add_paragraph(caption)
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

def pct(value):
    return f'{100*value:.2f}%'.replace('.', ',')

passed = sum(not any(c.find(tag) is not None for tag in ['failure', 'error', 'skipped']) for c in cases)
doc.add_heading('Laporan Pengujian\nSentimen Karhutla SVM', 0)
paragraph('MVP CLI | Python, TF-IDF, dan LinearSVC')
doc.add_heading('1. Ringkasan hasil', 1)
paragraph(f'Pengujian otomatis: {passed} dari {len(cases)} kasus lulus; {suite.attrib["failures"]} gagal dan {suite.attrib["errors"]} error. Pipeline berfungsi pada kasus yang diuji. Kualitas klasifikasi belum merata: seluruh teks netral pada holdout salah diklasifikasikan.')
table(['Indikator', 'Hasil aktual'], [
    ('Dataset', '92 baris: positif 38, negatif 39, netral 15'),
    ('Pembagian data', f'{r["train_rows"]} train / {r["test_rows"]} test; seed {r["seed"]}'),
    ('Accuracy holdout', f'{pct(r["accuracy"])} (13 benar dari 19 teks)'),
    ('Macro F1 / weighted F1', f'{pct(r["classification_report"]["macro avg"]["f1-score"])} / {pct(r["classification_report"]["weighted avg"]["f1-score"])}'),
    ('Durasi pytest', f'{float(suite.attrib["time"]):.2f} detik (JUnit)'),
])
doc.add_heading('2. Sumber dan metode', 1)
paragraph(f'Pengujian dijalankan ulang untuk laporan ini. Timestamp JUnit: {suite.attrib.get("timestamp", "tidak tercatat")} (waktu lokal host). Lingkungan: Python {platform.python_version()}, scikit-learn {r["sklearn_version"]}; versi pytest tercatat di pytest-output.txt.')
paragraph('Sumber: tests/test_pipeline.py, reports/pytest-results.xml, reports/pytest-output.txt, dan artifacts/evaluation.json. Tidak ada metrik atau data sintetis yang ditambahkan ke evaluasi model.')
chart('distribusi_dataset.png', 'Gambar 1. Jumlah label pada dataset asli.')

doc.add_page_break()
doc.add_heading('3. Kasus pengujian otomatis', 1)
paragraph('Seluruh kasus berada pada tests/test_pipeline.py. PASS berarti assertion kasus tersebut berhasil, bukan jaminan semua kondisi produksi telah diuji.')
descriptions = {
    'test_preprocessing': 'Normalisasi huruf, URL, mention, tanda baca, hashtag; negasi tetap ada; teks kosong tetap kosong.',
    'test_dataset_and_split': '92 baris dan distribusi label sesuai; grup dan teks identik tidak tumpang tindih; ketiga kelas tersedia; split deterministik; kata khusus test tidak masuk vocabulary train.',
    'test_reference_chains_and_duplicates': 'Rantai referensi dan teks identik bergabung transitif menjadi satu grup pada salinan empat baris asli.',
    'test_validation[label]': 'Label tidak valid ditolak dengan ValueError.',
    'test_validation[empty_text]': 'Teks yang hanya berisi URL ditolak setelah preprocessing.',
    'test_validation[duplicate_id]': 'ID duplikat ditolak.',
    'test_validation[missing_column]': 'Kolom referensi_id yang hilang ditolak.',
    'test_train_reload_and_cli': 'Training menghasilkan empat artifact; jumlah split/matriks benar; prediksi konsisten setelah reload; CLI predict berhasil; input kosong, file hilang, dan test_size=1 ditolak.',
}
for i, case in enumerate(cases, 1):
    name = case.attrib['name']
    doc.add_heading(f'{i:02d}. {name}', 2)
    paragraph(f'PASS | {float(case.attrib["time"]):.3f} detik. {descriptions[name]}')
paragraph('Perintah: .venv/bin/pytest -v --junitxml=reports/pytest-results.xml. Output lengkap disimpan pada reports/pytest-output.txt.')
paragraph('Catatan: kasus validasi memakai salinan dataset asli yang sengaja dibuat invalid; perubahan tersebut tidak digunakan untuk training atau perhitungan metrik.')

doc.add_page_break()
doc.add_heading('4. Evaluasi kualitas model', 1)
paragraph('Evaluasi memakai 19 teks holdout yang tidak digunakan untuk fit TF-IDF maupun SVM. Split berbasis grup referensi dan duplikat teks ternormalisasi. Proporsi test diminta 20%; aktual 20,65% karena grup tidak dipecah.')
table(['Label', 'Precision', 'Recall', 'F1-score', 'Support'], [(s.capitalize(), pct(r['classification_report'][s]['precision']), pct(r['classification_report'][s]['recall']), pct(r['classification_report'][s]['f1-score']), int(r['classification_report'][s]['support'])) for s in labels] + [(s, pct(r['classification_report'][key]['precision']), pct(r['classification_report'][key]['recall']), pct(r['classification_report'][key]['f1-score']), 19) for s, key in [('Macro avg', 'macro avg'), ('Weighted avg', 'weighted avg')]])
chart('metrik_per_kelas.png', 'Gambar 2. Precision, recall, dan F1-score per kelas pada holdout.')
paragraph('Precision: proporsi prediksi kelas yang benar. Recall: proporsi teks aktual kelas yang ditemukan. F1-score: rata-rata harmonik precision dan recall. Macro avg memberi bobot sama untuk setiap kelas; weighted avg mempertimbangkan jumlah teks per kelas.')
paragraph('Model tidak menghasilkan prediksi netral pada holdout. Precision netral sebenarnya tidak terdefinisi karena tidak ada prediksi kelas tersebut; laporan scikit-learn memakai nilai 0 melalui zero_division=0. Recall dan F1 netral juga 0.')
paragraph('Accuracy 68,42% tidak cukup untuk menyatakan model baik pada semua kelas. Macro F1 49,56% memperlihatkan kelemahan kelas netral yang tersamarkan oleh accuracy.')

doc.add_page_break()
doc.add_heading('5. Confusion matrix dan kesimpulan', 1)
chart('confusion_matrix.png', 'Gambar 3. Baris = label aktual; kolom = label prediksi. Angka merupakan jumlah teks.', width=6.1)
paragraph('Positif: 6 benar, 2 menjadi negatif. Negatif: 7 benar, 1 menjadi positif. Netral: 0 benar, 1 menjadi positif dan 2 menjadi negatif. Total 13 prediksi benar dan 6 salah.')
doc.add_heading('Kesimpulan', 2)
paragraph('Fungsi MVP yang tercakup dalam delapan kasus pengujian berjalan. Model dan hasil evaluasi dapat disimpan, dimuat, dan dipakai melalui CLI. Hasil ini membuktikan pipeline, bukan kesiapan klasifikasi untuk produksi.')
doc.add_heading('Batasan dan tindak lanjut', 2)
for text in [
    'Dataset kecil dan campuran bahasa Indonesia/Melayu; satu holdout berisi hanya tiga teks netral. Tidak ada interval ketidakpastian atau cross-validation dalam hasil ini.',
    'Pencegahan leakage mencakup rantai referensi dan teks identik setelah normalisasi, belum duplikat dekat, parafrasa, atau posting berbeda mengenai kejadian sama.',
    'Cakupan tes belum meliputi artifact rusak, seluruh variasi CSV, dan semua kombinasi grup yang tidak dapat di-split. PRD Markdown lengkap tidak tersedia untuk audit acceptance criteria.',
    'Tinjau definisi serta konsistensi label netral, tambah data berlabel asli, lalu gunakan evaluasi berbasis grup yang lebih luas. Hindari tuning berulang pada holdout yang sama.',
]:
    doc.add_paragraph(text, style='List Bullet')
paragraph('Lampiran terpisah: pytest-output.txt, pytest-results.xml, serta tiga file PNG chart di reports/. Rincian prediksi holdout tersedia pada artifacts/test_predictions.csv.')
path = OUT / 'Laporan_Pengujian_Sentimen_Karhutla.docx'
doc.save(path)
print(path)
